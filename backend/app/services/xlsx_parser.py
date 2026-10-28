from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.services.normalize import norm_key

SYSTEM_FIELDS = [
    ("nome", "Nome do Cliente"),
    ("endereco", "Endereço"),
    ("telefone", "Telefone / Contato"),
    ("rota", "Rota / Zona"),
    ("dia", "Dia da Semana"),
    ("ultima_visita", "Última Visita"),
    ("cidade", "Cidade"),
    ("bairro", "Bairro"),
    ("ignorar", "Ignorar coluna"),
]

_SUGGEST = {
    "razaosocial": "nome",
    "nome": "nome",
    "cliente": "nome",
    "nomefantasia": "nome",
    "endereco": "endereco",
    "end": "endereco",
    "contato": "telefone",
    "telefone": "telefone",
    "fone": "telefone",
    "zona": "rota",
    "rota": "rota",
    "zonarota": "rota",
    "diasemana": "dia",
    "dia": "dia",
    "ultvisita": "ultima_visita",
    "ultimavisita": "ultima_visita",
    "cidade": "cidade",
    "bairro": "bairro",
}


def detect_header_row(rows: list[tuple]) -> int:
    for i, row in enumerate(rows, start=1):
        strings = sum(1 for c in row if isinstance(c, str) and c.strip())
        if strings >= 3:
            return i
    raise ValueError("Não achei a linha de títulos na planilha.")


def parse_workbook(content: bytes, preview_rows: int = 10) -> dict[str, Any]:
    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    ws = wb.active
    raw_rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not raw_rows:
        raise ValueError("Planilha vazia.")

    header_idx = detect_header_row(raw_rows)
    header_row = raw_rows[header_idx - 1]
    headers: list[str] = []
    col_indexes: list[int] = []
    for i, cell in enumerate(header_row):
        if cell is None or str(cell).strip() == "":
            continue
        headers.append(str(cell).strip())
        col_indexes.append(i)

    suggested = {}
    for h in headers:
        suggested[h] = _SUGGEST.get(norm_key(h), "ignorar")

    data_rows = []
    for row in raw_rows[header_idx:]:
        values = [row[i] if i < len(row) else None for i in col_indexes]
        if not any(v is not None and str(v).strip() for v in values):
            continue
        data_rows.append(values)

    preview = []
    for values in data_rows[:preview_rows]:
        preview.append({headers[i]: values[i] for i in range(len(headers))})

    signature = "|".join(norm_key(h) for h in headers)
    return {
        "header_row": header_idx,
        "headers": headers,
        "suggested_mapping": suggested,
        "preview": preview,
        "total_rows": len(data_rows),
        "signature": signature,
        "rows": [
            {headers[i]: values[i] for i in range(len(headers))} for values in data_rows
        ],
    }
