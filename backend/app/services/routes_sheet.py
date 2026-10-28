from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

from app.config import get_settings
from app.services.days import DaySpec, parse_day_spec, week_dates
from app.services.normalize import norm_place


@dataclass
class CityRoute:
    cidade: str
    cidade_norm: str
    rota: str
    rota_id: str
    saida_carga: str
    enviar_cards: str
    card_days: DaySpec
    carga_days: DaySpec


@dataclass
class NeighborhoodRoute:
    bairro: str
    bairro_norm: str
    rota: str
    rota_id: str


@dataclass
class RouteCatalog:
    cities: list[CityRoute] = field(default_factory=list)
    neighborhoods: list[NeighborhoodRoute] = field(default_factory=list)
    by_city: dict[str, CityRoute] = field(default_factory=dict)
    by_bairro_vdc: dict[str, NeighborhoodRoute] = field(default_factory=dict)
    routes: dict[str, dict] = field(default_factory=dict)


def slug_rota(name: str) -> str:
    return norm_place(name).replace(" ", "-")


def load_catalog(path: str | None = None) -> RouteCatalog:
    xlsx = Path(path or get_settings().rotas_xlsx)
    wb = load_workbook(xlsx, data_only=True, read_only=True)
    catalog = RouteCatalog()

    ws = wb[wb.sheetnames[0]]
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i == 1 or not row or not row[0]:
            continue
        cidade = str(row[0]).strip()
        rota = str(row[1] or "").strip()
        saida = str(row[2] or "").strip()
        cards = str(row[3] or "").strip()
        item = CityRoute(
            cidade=cidade,
            cidade_norm=norm_place(cidade),
            rota=rota,
            rota_id=slug_rota(rota),
            saida_carga=saida,
            enviar_cards=cards,
            card_days=parse_day_spec(cards),
            carga_days=parse_day_spec(saida),
        )
        catalog.cities.append(item)
        catalog.by_city[item.cidade_norm] = item
        bucket = catalog.routes.setdefault(
            item.rota_id,
            {
                "id": item.rota_id,
                "nome": rota,
                "saida_carga": saida,
                "enviar_cards": cards,
                "cidades": [],
                "card_weekdays": item.card_days.weekdays,
                "carga_weekdays": item.carga_days.weekdays,
                "biweekly": item.card_days.biweekly,
            },
        )
        bucket["cidades"].append(cidade)

    if len(wb.sheetnames) > 1:
        ws2 = wb[wb.sheetnames[1]]
        for i, row in enumerate(ws2.iter_rows(values_only=True), start=1):
            if i == 1 or not row or not row[1]:
                continue
            bairro = str(row[1]).strip()
            rota = str(row[2] or "").strip()
            item = NeighborhoodRoute(
                bairro=bairro,
                bairro_norm=norm_place(bairro),
                rota=rota,
                rota_id=slug_rota(rota),
            )
            catalog.neighborhoods.append(item)
            catalog.by_bairro_vdc[item.bairro_norm] = item

    wb.close()
    return catalog


@lru_cache(maxsize=1)
def catalog() -> RouteCatalog:
    return load_catalog()


def resolve_route(cidade: str | None, bairro: str | None) -> dict | None:
    cat = catalog()
    city_norm = norm_place(cidade or "")
    # Divisão Leste/Oeste vale só para Vitória da Conquista.
    # Bairros como CENTRO existem em várias cidades e não podem vazar para VDC.
    if city_norm == "VITORIA DA CONQUISTA":
        bairro_hit = cat.by_bairro_vdc.get(norm_place(bairro or ""))
        vdc_meta = next((c for c in cat.cities if c.rota_id == (bairro_hit.rota_id if bairro_hit else "")), None)
        if not vdc_meta:
            vdc_meta = next((c for c in cat.cities if c.cidade_norm == "VITORIA DA CONQUISTA"), None)
        if bairro_hit:
            return {
                "id": bairro_hit.rota_id,
                "nome": bairro_hit.rota,
                "cidade": cidade,
                "bairro": bairro,
                "enviar_cards": vdc_meta.enviar_cards if vdc_meta else "Segunda e Quinta",
                "saida_carga": vdc_meta.saida_carga if vdc_meta else "Terça a Sexta",
                "origem": "bairro",
            }
        return None
    city = cat.by_city.get(city_norm)
    if city:
        return {
            "id": city.rota_id,
            "nome": city.rota,
            "cidade": city.cidade,
            "bairro": bairro,
            "enviar_cards": city.enviar_cards,
            "saida_carga": city.saida_carga,
            "origem": "cidade",
        }
    return None


def routes_for_day(day: date) -> list[dict]:
    cat = catalog()
    seen: dict[str, dict] = {}
    for city in cat.cities:
        if not city.card_days.matches(day):
            continue
        bucket = seen.setdefault(
            city.rota_id,
            {
                "id": city.rota_id,
                "nome": city.rota,
                "enviar_cards": city.enviar_cards,
                "saida_carga": city.saida_carga,
                "biweekly": city.card_days.biweekly,
                "cidades": [],
            },
        )
        if city.cidade not in bucket["cidades"]:
            bucket["cidades"].append(city.cidade)
    return list(seen.values())


def week_board(ref: date | None = None) -> list[dict]:
    board = []
    for day in week_dates(ref):
        routes = routes_for_day(day)
        board.append(
            {
                "date": day.isoformat(),
                "weekday": day.weekday(),
                "is_today": day == date.today(),
                "routes": routes,
            }
        )
    return board


def catalog_payload() -> dict:
    cat = catalog()
    return {
        "routes": list(cat.routes.values()),
        "cidades": [
            {
                "cidade": c.cidade,
                "rota": c.rota,
                "rota_id": c.rota_id,
                "enviar_cards": c.enviar_cards,
                "saida_carga": c.saida_carga,
                "card_days": c.card_days.weekdays,
            }
            for c in cat.cities
        ],
    }
